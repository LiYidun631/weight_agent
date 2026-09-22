from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path = Path('E:/llmpython/Weight_agent/output/Body270指标等级判定表.docx')
path.parent.mkdir(exist_ok=True)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el = borders.find(qn('w:'+edge))
        if el is None: el=OxmlElement('w:'+edge); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:color'),'D9D9D9')

def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for c,h in zip(t.rows[0].cells,headers): cell_text(c,h,True,'FFFFFF',8.5); shade(c,'2F5597')
    for j,row in enumerate(rows):
        for c,v in zip(t.add_row().cells,row): cell_text(c,v,size=8.5); shade(c,'F4F7FB' if j%2 else 'FFFFFF')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.5); sec.right_margin=Inches(.5)
doc.styles['Normal'].font.name='Microsoft YaHei'; doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); doc.styles['Normal'].font.size=Pt(9.5)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('八电极 Body270 指标等级判定表'); r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(19); r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('用于生成 AI 综合评估中的关键依据标签与说明文案'); r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(89,89,89)
p=doc.add_paragraph(); r=p.add_run('说明：'); r.bold=True; p.add_run('标准 Min/Max、身体类型、节段标准和控制量来自 BMH05108 八电极 Body270 算法输出。轻度、偏高、明显等严重程度属于业务展示规则，不代表协议官方等级。')
sections=[
('一、单指标判定表',['判定状态','判断条件','偏离率/分位数','展示标签','是否进入关键依据'],[
('数据无效','ERROR_TYPE != 0x00','-','数据异常','否'),('达标','Min ≤ value ≤ Max','-','达标项','可选'),('轻度偏高','value > Max','0%～5%','轻度偏高','是'),('偏高','value > Max','5%～20%','偏高','是'),('明显偏高','value > Max','>20%','明显偏高','是'),('轻度偏低','value < Min','0%～5%','轻度偏低','是'),('偏低','value < Min','5%～20%','偏低','是'),('明显偏低','value < Min','>20%','明显偏低','是'),('接近下限','在区间内，分位数 ≤20%','-','边界观察','可选'),('接近上限','在区间内，分位数 ≥80%','-','边界观察','可选'),('有利偏低','低于下限且低值有利','-','有利方向','否，不作为异常'),('无标准区间','缺少 Min 或 Max','-','暂不判定','否')]),
('二、指标分组判定表',['问题类别','主要指标','触发条件','关键依据文案'],[
('脂肪超标','体脂率、皮下脂肪率、体脂量、节段脂肪','体脂率或皮下脂肪率偏高；或节段脂肪至少 2 项为超标准','脂肪超标：体脂率……；皮下脂肪率……'),('肌肉不足','肌肉量、骨骼肌量、骨量、身体细胞量','肌肉量或骨骼肌量偏低；或核心肌肉指标至少 2 项偏低','肌肉不足：肌肉量……；骨骼肌量……'),('伴随偏低','水分量、蛋白质量、细胞内水量、细胞外水量','已判定肌肉不足，且伴随指标至少 2 项偏低','归入肌肉不足的伴随表现，不单独干预'),('正常体重肥胖','体重、BMI、体脂率、肌肉量','体重和 BMI 达标，同时体脂率偏高、肌肉量偏低','属于正常体重肥胖倾向'),('基础代谢偏低','基础代谢、肌肉量','基础代谢接近下限或偏低，且肌肉量偏低','基础代谢接近下限，可能与肌肉量不足相关'),('节段失衡','节段脂肪、节段肌肉','任一节段异常；至少 2 个节段异常时升级','存在局部脂肪/肌肉分布不均'),('达标项','BMI、内脏脂肪等级、肥胖度等','位于标准区间内','均在标准区间内，无需单独干预'),('有利方向','腰臀比等','低于下限，但低值为有利方向','低于标准下限，但属于有利方向')]),
('三、节段标准判定表',['协议值','含义','展示状态','建议处理'],[('0','低标准','偏低','计入节段异常数量'),('1','标准','达标','不计入异常'),('2','超标准','偏高','计入节段异常数量'),('0 项异常','节段分布正常','正常','无需单独说明'),('1 项异常','局部边界或轻度异常','局部观察','可作为补充依据'),('2 项异常','节段分布失衡','分布失衡','进入关键依据'),('≥3 项异常','明显节段失衡','明显失衡','作为重点问题')]),
('四、关键依据生成优先级表',['优先级','类型','处理方式','是否可覆盖低优先级结果'],[('1','数据异常','停止健康等级判定，仅提示测量异常','是'),('2','明显偏高/明显偏低','作为主要问题展示','是'),('3','偏高/偏低','作为主要问题展示','是'),('4','组合问题','生成“脂肪超标”“肌肉不足”等标签','是'),('5','伴随偏低','合并到主要问题，不单独干预','否'),('6','边界观察','作为补充说明，不升级为异常','否'),('7','达标项','汇总展示代表性指标','否'),('8','有利方向','单独说明，不列为异常','否')])]
for title,headers,rows in sections:
    h=doc.add_paragraph(); h.paragraph_format.space_before=Pt(7); h.paragraph_format.space_after=Pt(3); r=h.add_run(title); r.bold=True; r.font.size=Pt(12.5); r.font.color.rgb=RGBColor(31,78,121)
    table(doc,headers,rows)
h=doc.add_paragraph(); h.paragraph_format.space_before=Pt(5); r=h.add_run('计算公式'); r.bold=True; r.font.size=Pt(12.5); r.font.color.rgb=RGBColor(31,78,121)
for x in ['高于上限的偏离率 = (value - standard_max) / standard_max × 100%','低于下限的偏离率 = (standard_min - value) / standard_min × 100%','区间分位数 = (value - standard_min) / (standard_max - standard_min) × 100%']:
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1); p.add_run(x)
doc.save(path); print(path)
